"""
test_pipeline_e2e.py

End-to-end smoke tests. Run the real pipeline functions from main.py
against synthetic CSVs, using pytest's tmp_path for full filesystem
isolation — closest thing to a real user run, for both single-file
and batch modes.
"""

import openpyxl
import pandas as pd
import pytest

import main
import config


def _isolate(tmp_path, monkeypatch):
    """Redirect all output paths into the isolated tmp_path."""
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr(config, "CHARTS_DIR", tmp_path / "output" / "charts")
    monkeypatch.setattr(config, "LOG_PATH", tmp_path / "logs" / "data_quality.log")
    (tmp_path / "logs").mkdir(parents=True)


def test_run_single_produces_report(wind_df, tmp_path, monkeypatch):
    _isolate(tmp_path, monkeypatch)

    csv_path = tmp_path / "synthetic_wind.csv"
    wind_df.to_csv(csv_path, index=False)

    main.run_single(csv_path, tmp_path / "output", schema="wind", fmt="excel")

    reports = list((tmp_path / "output").glob("report_*.xlsx"))
    charts = list((tmp_path / "output" / "charts").glob("*.png"))

    assert len(reports) == 1
    assert len(charts) == 4
    assert (tmp_path / "logs" / "data_quality.log").exists()


def test_run_single_creates_missing_output_dir(wind_df, tmp_path, monkeypatch):
    """--output directory is created when it does not exist yet."""
    _isolate(tmp_path, monkeypatch)

    csv_path = tmp_path / "synthetic_wind.csv"
    wind_df.to_csv(csv_path, index=False)

    fresh_output = tmp_path / "deeply" / "nested" / "out"
    main.run_single(csv_path, fresh_output, schema="wind", fmt="excel")

    assert list(fresh_output.glob("report_*.xlsx")) != []


def test_run_single_html_report_is_self_contained(wind_df, tmp_path, monkeypatch):
    _isolate(tmp_path, monkeypatch)

    csv_path = tmp_path / "synthetic_wind.csv"
    wind_df.to_csv(csv_path, index=False)

    main.run_single(csv_path, tmp_path / "output", schema="wind", fmt="html")

    reports = list((tmp_path / "output").glob("report_*.html"))
    assert len(reports) == 1
    html = reports[0].read_text(encoding="utf-8")
    assert 'src="data:image/png;base64,' in html
    assert "<h2>Data Quality</h2>" in html


# ── Batch mode (Unit 7 done-when checklist) ──────────────────────────


@pytest.fixture
def batch_folder(wind_df, tmp_path):
    """Folder with two valid wind CSVs and one malformed CSV."""
    folder = tmp_path / "batch_in"
    folder.mkdir()
    wind_df.to_csv(folder / "site_a.csv", index=False)
    wind_df.to_csv(folder / "site_b.csv", index=False)
    (folder / "broken.csv").write_text("ColX,ColY\n1,2\n3,4\n")
    return folder


def test_run_batch_processes_all_valid_files(
    batch_folder, tmp_path, monkeypatch
):
    """Unit 7: every .csv processed; malformed one skipped with warning."""
    _isolate(tmp_path, monkeypatch)
    output_dir = tmp_path / "output"

    main.run_batch(batch_folder, output_dir, schema="wind", fmt="excel")

    report_path = next(output_dir.glob("report_*.xlsx"))
    wb = openpyxl.load_workbook(report_path)

    # Consolidated Clean Data holds both sites' rows. _write_dataframe
    # emits the unnamed index as column 1 plus a blank row after the
    # header, so data starts at row 3 with source_file in column 2.
    ws = wb["Clean Data"]
    site_values = {
        ws.cell(row=r, column=2).value for r in range(3, ws.max_row + 1)
    }
    assert site_values == {"site_a.csv", "site_b.csv"}


def test_run_batch_report_contains_all_run_ids(
    batch_folder, tmp_path, monkeypatch
):
    """Unit 7 fix: every file's run_id appears in the Data Quality Log."""
    _isolate(tmp_path, monkeypatch)
    output_dir = tmp_path / "output"

    main.run_batch(batch_folder, output_dir, schema="wind", fmt="excel")

    # Two successful files → exactly two run_ids logged
    import csv
    with open(config.LOG_PATH, newline="", encoding="utf-8") as f:
        logged_ids = {r["run_id"] for r in csv.DictReader(f)}
    assert len(logged_ids) == 2

    report_path = next(output_dir.glob("report_*.xlsx"))
    wb = openpyxl.load_workbook(report_path)
    ws = wb["Data Quality Log"]
    ids_in_sheet = {
        ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)
    }
    assert ids_in_sheet == logged_ids
